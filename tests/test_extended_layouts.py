from pptx import Presentation

from ppt_generate import DeckPipeline


def test_editable_kpi_and_timeline(tmp_path):
    deck = {
        "schema_version": "1.0",
        "title": "Report",
        "slides": [
            {
                "layout": "kpi",
                "title": "Operating metrics",
                "metrics": [
                    {"label": "Revenue", "value": "150 CNY", "detail": "2026 Q3"},
                    {"label": "Customers", "value": "300", "detail": "Paying accounts"},
                ],
            },
            {
                "layout": "timeline",
                "title": "Delivery milestones",
                "events": [
                    {
                        "date": "Q4",
                        "heading": "Acquisition",
                        "detail": "Sales diversifies customer acquisition",
                    },
                    {
                        "date": "Q4",
                        "heading": "Delivery",
                        "detail": "Operations adds weekly checkpoints",
                    },
                ],
            },
        ],
    }
    result = DeckPipeline(tmp_path).render(deck)
    presentation = Presentation(result["exported_file"])
    assert all(shape.has_text_frame for slide in presentation.slides for shape in slide.shapes)
    text = " ".join(shape.text for slide in presentation.slides for shape in slide.shapes)
    assert "150 CNY" in text and "weekly checkpoints" in text
