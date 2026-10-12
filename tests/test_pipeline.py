import asyncio
import copy
import json
from pathlib import Path

import pytest
from fastmcp import Client
from pptx import Presentation

from ppt_generate.core.pipeline import DeckPipeline
from ppt_generate.mcp.server import PPTGenerateServer


@pytest.fixture
def deck():
    return json.loads(
        (Path(__file__).resolve().parents[1] / "examples/quarterly_report.json").read_text(
            encoding="utf-8"
        )
    )


def test_native_objects_and_isolated_jobs(tmp_path, deck):
    results = [DeckPipeline(tmp_path).render(deck) for _ in range(3)]
    assert len({r["exported_file"] for r in results}) == 3
    for result in results:
        prs = Presentation(result["exported_file"])
        assert len(prs.slides) == 8
        assert sum(s.has_chart for sl in prs.slides for s in sl.shapes) == 2
        assert sum(s.has_table for sl in prs.slides for s in sl.shapes) == 2
        assert prs.slides[4].shapes[2].chart.series[0].values == (100.0, 125.0, 150.0)
        assert any("模拟数据" in s.text for s in prs.slides[0].shapes if s.has_text_frame)


def test_revision_and_rollback(tmp_path, deck):
    pipeline = DeckPipeline(tmp_path)
    original = pipeline.render(deck)
    replacement = copy.deepcopy(deck["slides"][2])
    replacement["title"] = "修订经营摘要"
    revised = pipeline.revise(original["job_id"], 3, replacement)
    new = json.loads(Path(revised["deck_file"]).read_text(encoding="utf-8"))
    assert new["slides"][:2] == deck["slides"][:2]
    assert new["slides"][3:] == deck["slides"][3:]
    assert new["slides"][2] == replacement
    manifest = tmp_path / original["job_id"] / "manifest.json"
    before = manifest.read_bytes()
    with pytest.raises(ValueError):
        pipeline.revise(
            original["job_id"], 3, {"layout": "bullets", "title": "Bad", "bullets": ["x" * 111]}
        )
    assert manifest.read_bytes() == before
    assert Path(original["exported_file"]).exists()


@pytest.mark.parametrize(
    "mutate",
    [
        lambda d: d["slides"][4]["series"][0].update(values=[1]),
        lambda d: d["slides"][4]["series"][0].update(values=[1, float("nan"), 3]),
        lambda d: d["slides"][3]["rows"].append(["wrong width"]),
        lambda d: d["slides"][2].update(bullets=["中" * 110] * 6),
    ],
)
def test_rejection_without_publishing(tmp_path, deck, mutate):
    mutate(deck)
    with pytest.raises(ValueError):
        DeckPipeline(tmp_path).render(deck)
    assert not list(tmp_path.rglob("manifest.json"))
    assert not list(tmp_path.rglob("presentation.pptx"))


def test_path_traversal(tmp_path):
    with pytest.raises(ValueError):
        DeckPipeline(tmp_path).revise("../outside", 1, {})


def test_mcp_contract(tmp_path, deck):
    async def run():
        async with Client(PPTGenerateServer(tmp_path).app) as client:
            assert {t.name for t in await client.list_tools()} == {
                "create_ppt",
                "render_deck",
                "revise_slide",
                "get_deck_schema",
                "get_job",
                "read_document",
                "preview_ppt",
                "repair_ppt",
                "bind_evidence",
                "record_review",
                "submit_ppt",
                "get_task",
                "cancel_task",
                "cleanup_history",
            }
            result = await client.call_tool("render_deck", {"deck": deck})
            assert result.data["success"] is True
            bad = await client.call_tool("create_ppt", {"topic": "Demo", "export_format": "pdf"})
            assert bad.data["success"] is False

    asyncio.run(run())
