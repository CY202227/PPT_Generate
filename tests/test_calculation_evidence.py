import copy
import json
from pathlib import Path

import pytest

from ppt_generate import DeckPipeline
from ppt_generate.core.evidence import audit_bindings


def case(operation="percentage_change", value=25, unit="%", operands=(100, 125), precision=2):
    quotes = [f"2026 Q{i + 1} Revenue {number} CNY" for i, number in enumerate(operands)]
    source = "\n".join(quotes)
    inputs = [
        {
            "value": number,
            "start": source.index(quote),
            "quote": quote,
            "category": f"Q{i + 1}",
            "unit": "CNY",
            "period": "2026",
        }
        for i, (number, quote) in enumerate(zip(operands, quotes))
    ]
    deck = {
        "schema_version": "1.0",
        "title": "Report",
        "slides": [
            {
                "layout": "chart",
                "title": "Calculated comparison",
                "source": "Synthetic data",
                "unit": unit,
                "period": "2026",
                "categories": ["Comparison"],
                "series": [{"name": "Revenue", "values": [value]}],
            }
        ],
    }
    bindings = [
        {
            "path": "/slides/0/series/0/values/0",
            "derivation": {"operation": operation, "precision": precision, "inputs": inputs},
        }
    ]
    return deck, source, bindings


@pytest.mark.parametrize(
    "operation,value,unit,operands",
    [
        ("sum", 225, "CNY", (100, 125)),
        ("difference", -25, "CNY", (100, 125)),
        ("ratio", 0.8, "ratio", (100, 125)),
        ("percentage", 80, "%", (100, 125)),
        ("percentage_change", 25, "%", (100, 125)),
    ],
)
def test_supported_calculations(operation, value, unit, operands):
    deck, source, bindings = case(operation, value, unit, operands)
    report = audit_bindings(deck, source, bindings)
    assert report["status"] == "context_checks_passed"
    assert report["references"][0]["status"] == "calculation_checks_passed"
    assert report["factual_accuracy"] == "not_verified"


def test_arithmetic_failure_does_not_get_numeric_exemption(tmp_path):
    deck, source, bindings = case(value=30)
    pipeline = DeckPipeline(tmp_path)
    original = pipeline.render(deck, source_material=source)
    result = pipeline.bind_evidence(original["job_id"], bindings, original["revision"])
    assert result["delivery"]["blocked_by"] == ["evidence", "provenance"]
    report = json.loads(Path(result["quality_report"]).read_text(encoding="utf-8"))
    assert report["provenance"]["issues"][0]["code"] == "calculation_result_mismatch"


def test_valid_derivation_only_exempts_its_chart_path_and_resets_review(tmp_path):
    deck, source, bindings = case()
    pipeline = DeckPipeline(tmp_path)
    original = pipeline.render(deck, source_material=source)
    assert original["delivery"]["blocked_by"] == ["evidence"]
    good = pipeline.bind_evidence(original["job_id"], bindings, original["revision"])
    assert good["delivery"]["blocked_by"] == []
    assert not good["delivery"]["ready"]
    pipeline.record_review(
        original["job_id"], "factual_review", "approved", "Test reviewer", original["revision"]
    )
    rebound = pipeline.bind_evidence(original["job_id"], bindings, original["revision"])
    assert "factual_review" in rebound["delivery"]["pending"]
    new_deck = copy.deepcopy(deck)
    new_deck["slides"][0]["title"] = "Unsupported claim 25"
    changed = pipeline.revise(original["job_id"], 1, new_deck["slides"][0])
    rebound = pipeline.bind_evidence(changed["job_id"], bindings, changed["revision"])
    assert "evidence" in rebound["delivery"]["blocked_by"]


@pytest.mark.parametrize(
    "mutation,code",
    [
        (
            lambda b: b[0]["derivation"]["inputs"][0].update(quote="Private wrong source"),
            "operand_quote_mismatch",
        ),
        (lambda b: b[0]["derivation"]["inputs"][0].update(value=99), "operand_value_missing"),
        (
            lambda b: b[0]["derivation"]["inputs"][0].update(unit="USD"),
            "calculation_input_unit_mismatch",
        ),
        (
            lambda b: b[0]["derivation"]["inputs"][0].update(period="2025"),
            "calculation_period_mismatch",
        ),
    ],
)
def test_bad_operands_block_context(mutation, code):
    deck, source, bindings = case()
    mutation(bindings)
    report = audit_bindings(deck, source, bindings)
    assert code in {issue["code"] for issue in report["issues"]}
    assert "Private wrong source" not in json.dumps(report)


def test_zero_denominator_and_half_up_rounding():
    deck, source, bindings = case(operands=(0, 100))
    report = audit_bindings(deck, source, bindings)
    assert "calculation_zero_denominator" in {issue["code"] for issue in report["issues"]}
    deck, source, bindings = case("percentage", 12.5, "%", (1, 8), 1)
    assert audit_bindings(deck, source, bindings)["status"] == "context_checks_passed"
    deck, source, bindings = case("ratio", 0.13, "ratio", (1, 8), 2)
    assert (
        audit_bindings(deck, source, bindings)["references"][0]["calculation"]["computed_value"]
        == "0.13"
    )


def test_no_arbitrary_expression_or_nonfinite_operand():
    deck, source, bindings = case()
    bindings[0]["derivation"]["operation"] = "__import__('os').system('unexpected')"
    with pytest.raises(ValueError, match="Unsupported"):
        audit_bindings(deck, source, bindings)

    deck, source, bindings = case()
    bindings[0]["derivation"]["inputs"][0]["value"] = "1e999999"
    with pytest.raises(ValueError, match="range"):
        audit_bindings(deck, source, bindings)


def test_schema_and_cross_period_calculation():
    from jsonschema import Draft202012Validator

    from ppt_generate.core.evidence import get_evidence_schema

    schema = get_evidence_schema()
    Draft202012Validator.check_schema(schema)
    deck, source, bindings = case()
    Draft202012Validator(schema).validate(bindings)
    source = source.replace("2026 Q1", "2025 Q1")
    operand = bindings[0]["derivation"]["inputs"][0]
    operand.update(quote=operand["quote"].replace("2026", "2025"), period="2025")
    deck["slides"][0]["period"] = "2025–2026"
    assert audit_bindings(deck, source, bindings)["status"] == "context_checks_passed"


def test_repair_keeps_calculation_evidence_and_remaps_paths(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Pt

    deck, source, bindings = case()
    pipeline = DeckPipeline(tmp_path)
    saved = pipeline.render(deck, source_material=source)
    pipeline.bind_evidence(saved["job_id"], bindings, saved["revision"])
    presentation = Presentation(saved["exported_file"])
    title = presentation.slides[0].shapes[0]
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.size = Pt(60)
    presentation.save(saved["exported_file"])
    reports = [
        [
            {
                "page": 1,
                "object": title.name,
                "code": "text_overflow",
                "bound_height_pt": 120,
                "available_height_pt": 60,
            }
        ],
        [],
    ]

    def preview(pptx, output, backend):
        issues = reports.pop(0)
        report = {
            "backend": "powerpoint",
            "status": "issues_found" if issues else "passed",
            "issues": issues,
            "pages": [],
        }
        (output / "report.json").write_text(json.dumps(report), encoding="utf-8")
        return report

    monkeypatch.setattr("ppt_generate.core.preview.render_preview", preview)
    pipeline.preview(saved["job_id"])
    repaired = pipeline.repair_preview(saved["job_id"])
    assert Path(repaired["bindings_report"]).is_file()
    report = json.loads(Path(repaired["quality_report"]).read_text(encoding="utf-8"))
    assert report["provenance"]["references"][0]["status"] == "calculation_checks_passed"
    assert not repaired["delivery"]["blocked_by"]
    assert not repaired["delivery"]["ready"]
