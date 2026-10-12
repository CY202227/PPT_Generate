import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

from ppt_generate import DeckPipeline
from ppt_generate.core.pipeline import PlanningFailed


def fixture(tmp_path):
    pipeline = DeckPipeline(tmp_path)
    deck = {
        "schema_version": "1.0",
        "title": "Report",
        "slides": [
            {"layout": "cover", "title": "Report"},
            {"layout": "bullets", "title": "Results", "bullets": ["Revenue 150"]},
        ],
    }
    saved = pipeline.render(deck, source_material="Revenue 150")
    presentation = Presentation(saved["exported_file"])
    shape = next(
        s for s in presentation.slides[1].shapes if s.has_text_frame and s.text == "Results"
    )
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.size = Pt(60)
    presentation.save(saved["exported_file"])
    issue = {
        "page": 2,
        "object": shape.name,
        "code": "text_overflow",
        "bound_height_pt": 72,
        "available_height_pt": 40,
        "bound_width_pt": 100,
        "available_width_pt": 800,
    }
    return pipeline, saved, issue


def engine(monkeypatch, reports, callback=None):
    def fake(pptx, output, backend):
        issues = reports.pop(0)
        if callback:
            callback()
        report = {
            "backend": "powerpoint",
            "status": "issues_found" if issues else "passed",
            "pages": [],
            "issues": issues,
        }
        (output / "report.json").write_text(json.dumps(report), encoding="utf-8")
        return report

    monkeypatch.setattr("ppt_generate.core.preview.render_preview", fake)


def test_only_flagged_shape_changes_and_clean_candidate_promoted(tmp_path, monkeypatch):
    pipeline, saved, issue = fixture(tmp_path)
    engine(monkeypatch, [[issue], []])
    preview = pipeline.preview(saved["job_id"])
    before = Presentation(saved["exported_file"])
    result = pipeline.repair_preview(saved["job_id"], preview["revision"])
    after = Presentation(result["exported_file"])
    assert before.slides[0]._element.xml == after.slides[0]._element.xml
    assert [s.text for s in before.slides[1].shapes if s.has_text_frame] == [
        s.text for s in after.slides[1].shapes if s.has_text_frame
    ]
    assert pipeline.get_job(saved["job_id"])["deck"] == json.loads(
        Path(saved["deck_file"]).read_text()
    )
    assert result["revision"] != saved["revision"]
    assert result["delivery"]["status"] == "needs_review"
    assert Path(result["preview_report"]).is_file()
    assert json.loads(Path(result["repair_report"]).read_text())["status"] == "accepted"


def test_unsupported_defect_and_engine_failure_preserve_manifest(tmp_path, monkeypatch):
    pipeline, saved, issue = fixture(tmp_path)
    issue["code"] = "text_overlap"
    engine(monkeypatch, [[issue]])
    pipeline.preview(saved["job_id"])
    manifest = tmp_path / saved["job_id"] / "manifest.json"
    original = manifest.read_bytes()
    with pytest.raises(PlanningFailed):
        pipeline.repair_preview(saved["job_id"])
    assert manifest.read_bytes() == original
    issue["code"] = "text_overflow"
    engine(monkeypatch, [[issue]])
    pipeline.preview(saved["job_id"])
    original = manifest.read_bytes()

    def unavailable(*args):
        raise RuntimeError("engine failure")

    monkeypatch.setattr("ppt_generate.core.preview.render_preview", unavailable)
    with pytest.raises(RuntimeError):
        pipeline.repair_preview(saved["job_id"])
    assert manifest.read_bytes() == original


def test_concurrent_edit_wins_over_repair(tmp_path, monkeypatch):
    pipeline, saved, issue = fixture(tmp_path)
    engine(monkeypatch, [[issue]])
    pipeline.preview(saved["job_id"])

    def edit():
        pipeline.revise(saved["job_id"], 1, {"layout": "cover", "title": "New cover"})

    engine(monkeypatch, [[]], edit)
    with pytest.raises(ValueError, match="revision conflict"):
        pipeline.repair_preview(saved["job_id"])
    assert pipeline.get_job(saved["job_id"])["deck"]["slides"][0]["title"] == "New cover"


def test_two_attempt_limit_and_font_floor(tmp_path, monkeypatch):
    pipeline, saved, issue = fixture(tmp_path)
    engine(monkeypatch, [[issue], [issue], [issue]])
    pipeline.preview(saved["job_id"])
    with pytest.raises(PlanningFailed) as error:
        pipeline.repair_preview(saved["job_id"])
    assert len(error.value.diagnostics["attempts"]) == 2
    assert all(
        change["minimum_font_pt"] >= 24
        for attempt in error.value.diagnostics["attempts"]
        for change in attempt["changes"]
    )
    assert pipeline.get_job(saved["job_id"])["manifest"]["revision"] == saved["revision"]


def test_review_during_repair_is_not_overwritten(tmp_path, monkeypatch):
    pipeline, saved, issue = fixture(tmp_path)
    engine(monkeypatch, [[issue]])
    pipeline.preview(saved["job_id"])

    def review():
        pipeline.record_review(
            saved["job_id"], "factual_review", "approved", "Test reviewer", saved["revision"]
        )

    engine(monkeypatch, [[]], review)
    with pytest.raises(ValueError, match="revision conflict"):
        pipeline.repair_preview(saved["job_id"])
    quality = json.loads(
        Path(pipeline.get_job(saved["job_id"])["manifest"]["quality_report"]).read_text()
    )
    assert quality["reviews"]["factual_review"]["decision"] == "approved"
