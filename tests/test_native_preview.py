import asyncio
import copy
import json
from pathlib import Path

from fastmcp import Client
from pptx import Presentation
from pptx.util import Inches

from ppt_generate.core.native_preview import resolve_pages
from ppt_generate.core.pipeline import DeckPipeline
from ppt_generate.core.preview import PreviewUnavailable
from ppt_generate.core.quality import delivery_gate
from ppt_generate.mcp.server import PPTGenerateServer


def sample():
    return json.loads(
        (Path(__file__).resolve().parents[1] / "examples/quarterly_report.json").read_text(
            encoding="utf-8"
        )
    )


def test_native_preview_without_external_engine(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Office subprocess must not run")

    monkeypatch.setattr("subprocess.run", forbidden)
    pipeline = DeckPipeline(tmp_path)
    deck = sample()
    deck["slides"][2]["title"] = '<script>alert("x")</script>'
    original = pipeline.render(deck)
    result = pipeline.preview(original["job_id"], "auto", original["revision"])
    report = json.loads(Path(result["preview_report"]).read_text(encoding="utf-8"))
    assert report["backend"] == "native"
    assert report["office_verification"] == "not_performed"
    assert len(report["pages"]) == 8
    html = Path(result["browser_preview"]).read_text(encoding="utf-8")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    images = list(Path(result["preview_report"]).parent.glob("slide-*.png"))
    assert len(images) == 8
    assert all(image.read_bytes().startswith(b"\x89PNG") for image in images)
    assert result["visual_review"] == "required"


def test_scene_uses_exported_geometry(tmp_path):
    result = DeckPipeline(tmp_path).render(sample())
    presentation = Presentation(result["exported_file"])
    presentation.slides[2].shapes[0].left = Inches(1.7)
    presentation.save(result["exported_file"])
    scene = resolve_pages(result["exported_file"])
    assert scene[2]["elements"][0]["x"] == 122.4
    charts = [item for page in scene for item in page["elements"] if item["kind"] == "chart"]
    assert len(charts) == 2
    assert charts[0]["series"][0]["values"] == [100, 125, 150]


def test_mcp_image_edit_and_conflict(tmp_path):
    async def run():
        server = PPTGenerateServer(tmp_path)
        original = server.pipeline.render(sample())
        async with Client(server.app) as client:
            args = {
                "job_id": original["job_id"],
                "page_number": 3,
                "base_revision": original["revision"],
            }
            read = await client.call_tool("get_slide", args)
            replacement = copy.deepcopy(read.data["slide"])
            replacement["title"] = "修订经营摘要"
            preview = await client.call_tool("preview_slide", args)
            assert any(item.type == "image" for item in preview.content)
            assert preview.data["revision"] == original["revision"]
            updated = await client.call_tool("update_slide", {**args, "slide": replacement})
            assert any(item.type == "image" for item in updated.content)
            assert updated.data["revision"] != original["revision"]
            assert Path(updated.data["exported_file"]).exists()
            saved = server.pipeline.get_job(original["job_id"])
            assert saved["deck"]["slides"][2] == replacement
            assert saved["deck"]["slides"][:2] == sample()["slides"][:2]
            assert saved["deck"]["slides"][3:] == sample()["slides"][3:]
            stale = await client.call_tool("update_slide", {**args, "slide": replacement})
            assert stale.data["error"]["code"] == "conflict"
            assert (
                server.pipeline.get_job(original["job_id"])["manifest"]["revision"]
                == updated.data["revision"]
            )

    asyncio.run(run())


def test_edit_keeps_export_when_preview_unavailable(tmp_path, monkeypatch):
    async def run():
        server = PPTGenerateServer(tmp_path)
        original = server.pipeline.render(sample())

        def unavailable(*args, **kwargs):
            raise PreviewUnavailable("Install preview extra")

        monkeypatch.setattr(server.pipeline, "preview_slide", unavailable)
        async with Client(server.app) as client:
            updated = await client.call_tool(
                "update_slide",
                {
                    "job_id": original["job_id"],
                    "page_number": 3,
                    "slide": sample()["slides"][2],
                    "base_revision": original["revision"],
                },
            )
            assert updated.data["success"] is True
            assert updated.data["preview_error"]["code"] == "preview_unavailable"
            assert Path(updated.data["exported_file"]).exists()
            assert updated.data["revision"] != original["revision"]

    asyncio.run(run())


def test_native_pass_does_not_hide_office_issues():
    quality = {
        "layout": {"status": "passed", "backend": "native", "issues": []},
        "office_layout": {
            "status": "issues_found",
            "backend": "powerpoint",
            "issues": [{"page": 1, "code": "text_overflow"}],
        },
    }
    assert delivery_gate(quality)["status"] == "blocked"
