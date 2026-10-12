from pptx import Presentation
from pptx.util import Inches

from ppt_generate import DeckPipeline
from ppt_generate.core.repair import repair_layout


def test_move_only_flagged_objects_and_preserve_text(tmp_path):
    saved = DeckPipeline(tmp_path).render(
        {
            "schema_version": "1.0",
            "title": "Report",
            "slides": [{"layout": "bullets", "title": "Results", "bullets": ["Revenue 150"]}],
        }
    )
    presentation = Presentation(saved["exported_file"])
    slide = presentation.slides[0]
    title = next(s for s in slide.shapes if s.has_text_frame and s.text == "Results")
    body = next(s for s in slide.shapes if s.has_text_frame and "Revenue" in s.text)
    body.top = title.top
    presentation.save(saved["exported_file"])
    title_xml = title._element.xml
    changes = repair_layout(
        saved["exported_file"],
        [{"page": 1, "code": "text_overlap", "objects": [title.name, body.name]}],
    )
    repaired = Presentation(saved["exported_file"])
    assert changes[0]["code"] == "separate_textboxes"
    assert repaired.slides[0].shapes[0]._element.xml == title_xml
    body = next(s for s in repaired.slides[0].shapes if s.name == body.name)
    assert "Revenue 150" in body.text
    body.left = -Inches(0.3)
    repaired.save(saved["exported_file"])
    changes = repair_layout(
        saved["exported_file"], [{"page": 1, "code": "object_out_of_bounds", "object": body.name}]
    )
    assert changes[0]["code"] == "move_inside_slide"
    assert (
        next(
            s for s in Presentation(saved["exported_file"]).slides[0].shapes if s.name == body.name
        ).left
        == 0
    )
