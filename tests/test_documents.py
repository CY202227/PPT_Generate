import pytest

from ppt_generate.core.documents import extract_document


def test_markdown_and_docx_reading_order(tmp_path):
    from docx import Document

    markdown = tmp_path / "source.md"
    markdown.write_text("# Revenue\n100", encoding="utf-8")
    assert extract_document(markdown)["text"] == "# Revenue\n100"
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "100"
    document.add_paragraph("After")
    source = tmp_path / "source.docx"
    document.save(source)
    assert extract_document(source)["text"] == "Before\n\nMetric | 100\n\nAfter"


def test_pdf_text_and_scanned_rejection(tmp_path):
    import fitz

    source = tmp_path / "source.pdf"
    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "Revenue 100")
        document.save(source)
    assert "Revenue 100" in extract_document(source)["text"]
    empty = tmp_path / "scanned.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(empty)
    with pytest.raises(ValueError, match="OCR"):
        extract_document(empty)


def test_document_limits_and_types(tmp_path):
    source = tmp_path / "large.md"
    source.write_text("x" * 100001, encoding="utf-8")
    with pytest.raises(ValueError, match="100,000"):
        extract_document(source)
    unsupported = tmp_path / "source.exe"
    unsupported.write_text("unsupported")
    with pytest.raises(ValueError, match="Supported inputs"):
        extract_document(unsupported)
