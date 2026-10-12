import json
from pathlib import Path

from pptx import Presentation

from ppt_generate import DeckPipeline


def test_english_split_and_negative_chart(tmp_path):
    source = Path(__file__).resolve().parents[1] / "examples/negative_report.json"
    deck = json.loads(source.read_text(encoding="utf-8"))
    saved = DeckPipeline(tmp_path).render(deck)
    prs = Presentation(saved["exported_file"])
    chart = next(shape.chart for shape in prs.slides[2].shapes if shape.has_chart)
    assert chart.series[0].values == (-30.0, -10.0, 10.0)
    assert chart.value_axis.minimum_scale is None
    assert chart.series[0].invert_if_negative is False
    assert chart.legend.include_in_layout is False
    assert any(
        "Distribute delivery" in shape.text
        for shape in prs.slides[4].shapes
        if shape.has_text_frame
    )


def test_mixed_language_subtitle_does_not_split_short_english_word(tmp_path):
    deck = {
        "schema_version": "1.0",
        "title": "Report",
        "slides": [
            {
                "layout": "cover",
                "title": "Report",
                "subtitle": "2026 Q3 管理层汇报；数据为模拟数据，来源：synthetic operating ledger",
            }
        ],
    }
    saved = DeckPipeline(tmp_path).render(deck)
    prs = Presentation(saved["exported_file"])
    subtitle = prs.slides[0].shapes[1].text
    assert "ledger" in subtitle
    assert "l\nedger" not in subtitle
