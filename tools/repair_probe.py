"""Opt-in Windows PowerPoint regression probe for content-preserving repair."""

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_generate import DeckPipeline


def main():
    root = Path("outputs/repair-probe").resolve()
    pipeline = DeckPipeline(root)
    deck = {
        "schema_version": "1.0",
        "title": "Repair verification",
        "slides": [
            {"layout": "cover", "title": "Repair verification", "subtitle": "Synthetic data"},
            {"layout": "bullets", "title": "Results", "bullets": ["Revenue 150", "Customers 300"]},
            {
                "layout": "bullets",
                "title": "Plan",
                "bullets": ["Owner: Operations", "Deadline: Q4"],
            },
        ],
    }
    deck["slides"].extend(
        [
            {
                "layout": "kpi",
                "title": "Operating snapshot",
                "metrics": [
                    {"label": "Revenue", "value": "150", "detail": "Synthetic report"},
                    {"label": "Customers", "value": "300", "detail": "Synthetic accounts"},
                    {"label": "Next period", "value": "Q4", "detail": "Operations plan"},
                ],
            },
            {
                "layout": "timeline",
                "title": "Delivery milestones",
                "events": [
                    {"date": "Q4", "heading": "Acquisition", "detail": "Sales owns acquisition"},
                    {"date": "Q4", "heading": "Delivery", "detail": "Operations owns checkpoints"},
                    {"date": "Q4", "heading": "Review", "detail": "Review measured results"},
                    {"date": "Q4", "heading": "Reporting", "detail": "Publish reviewed results"},
                ],
            },
        ]
    )
    original = pipeline.render(deck, source_material="Revenue 150. Customers 300.")
    presentation = Presentation(original["exported_file"])
    title = next(
        s for s in presentation.slides[1].shapes if s.has_text_frame and s.text == "Results"
    )
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(100)
    body = next(
        s for s in presentation.slides[1].shapes if s.has_text_frame and "Revenue" in s.text
    )
    body.top = title.top
    plan = next(s for s in presentation.slides[2].shapes if s.has_text_frame and "Owner" in s.text)
    plan.left = -Inches(0.3)
    presentation.save(original["exported_file"])
    before = pipeline.preview(original["job_id"], "powerpoint")
    assert before["layout_review"] == "issues_found"
    report = json.loads(Path(before["preview_report"]).read_text(encoding="utf-8-sig"))
    assert {"text_overflow", "text_overlap", "object_out_of_bounds"} <= {
        i["code"] for i in report["issues"]
    }
    repaired = pipeline.repair_preview(original["job_id"], before["revision"])
    assert repaired["layout_review"] == "passed"
    assert pipeline.get_job(original["job_id"])["deck"] == deck
    a, b = Presentation(original["exported_file"]), Presentation(repaired["exported_file"])
    assert a.slides[0]._element.xml == b.slides[0]._element.xml
    for index in (3, 4):
        assert a.slides[index]._element.xml == b.slides[index]._element.xml
    for index in (1, 2):
        assert [s.text for s in a.slides[index].shapes if s.has_text_frame] == [
            s.text for s in b.slides[index].shapes if s.has_text_frame
        ]
    result = {"original": before, "repaired": repaired, "content_preserved": True}
    (root / "result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps({"passed": True, "result_file": str(root / "result.json")}, indent=2))


if __name__ == "__main__":
    main()
