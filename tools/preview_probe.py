"""Opt-in real-engine check that defects are detected, not just clean files exported."""

import json
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_generate.core.preview import render_preview


def main():
    output = Path("tmp/preview-probes") / uuid.uuid4().hex
    output.mkdir(parents=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.15))
    first.text = "Oversized text"
    first.text_frame.paragraphs[0].font.size = Pt(42)
    second = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    second.text = "Overlapping text"
    second.text_frame.paragraphs[0].font.size = Pt(24)
    outside = slide.shapes.add_textbox(Inches(13), Inches(5), Inches(2), Inches(1))
    outside.text = "Outside"
    pptx = output / "defects.pptx"
    prs.save(pptx)
    report = render_preview(pptx, output / "previews", "powerpoint")
    codes = {issue["code"] for issue in report["issues"]}
    assert {"text_overflow", "text_overlap", "object_out_of_bounds"} <= codes, report
    print(
        json.dumps(
            {"detected": sorted(codes), "report": str(output / "previews/report.json")}, indent=2
        )
    )


if __name__ == "__main__":
    main()
