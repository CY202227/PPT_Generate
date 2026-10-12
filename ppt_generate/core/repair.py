"""Bounded, content-preserving repairs driven by measured PowerPoint overflow."""

import math

from pptx import Presentation
from pptx.util import Pt


def repair_overflow(path, issues):
    """Change only flagged text frames, never wording, data, or unflagged pages.

    This intentionally declines overlap/out-of-bounds and fonts below readable limits.
    A real preview must validate the resulting candidate before publication.
    """
    presentation = Presentation(path)
    changes = []
    seen = set()
    for issue in issues:
        if issue.get("code") != "text_overflow":
            continue
        page, name = issue.get("page"), issue.get("object")
        if (
            isinstance(page, bool)
            or not isinstance(page, int)
            or not 1 <= page <= len(presentation.slides)
        ):
            continue
        if not isinstance(name, str) or (page, name) in seen:
            continue
        seen.add((page, name))
        shape = next((s for s in presentation.slides[page - 1].shapes if s.name == name), None)
        frame = shape.text_frame if shape is not None and shape.has_text_frame else None
        if frame is None and ":" in name:
            shape_name, row, column = name.rsplit(":", 2)
            shape = next(
                (s for s in presentation.slides[page - 1].shapes if s.name == shape_name), None
            )
            if shape is not None and shape.has_table and row.isdigit() and column.isdigit():
                r, c = int(row), int(column)
                if 1 <= r <= len(shape.table.rows) and 1 <= c <= len(shape.table.columns):
                    frame = shape.table.cell(r - 1, c - 1).text_frame
        if frame is None:
            continue
        ratios = []
        for dimension in ("height", "width"):
            bound = issue.get(f"bound_{dimension}_pt")
            available = issue.get(f"available_{dimension}_pt")
            if (
                isinstance(bound, (int, float))
                and isinstance(available, (int, float))
                and math.isfinite(bound)
                and math.isfinite(available)
                and bound > 0
                and available > 0
            ):
                ratios.append(available / bound)
        if not ratios or min(ratios) >= 1:
            continue
        runs = [
            (run, run.font.size or paragraph.font.size)
            for paragraph in frame.paragraphs
            for run in paragraph.runs
        ]
        # Unknown inherited fonts cannot safely be shrunk without knowing their size.
        if not runs or any(size is None for _, size in runs):
            continue
        minimum = 24 if shape is not None and shape.top < Pt(80) else 18
        ratio = min(0.95, min(ratios) * 0.95)
        sizes = [max(minimum, math.floor(size.pt * ratio * 2) / 2) for _, size in runs]
        if any(size > old_size.pt for size, (_, old_size) in zip(sizes, runs)):
            continue
        if all(size == old_size.pt for size, (_, old_size) in zip(sizes, runs)):
            continue
        for (run, _), size in zip(runs, sizes):
            run.font.size = Pt(size)
        changes.append(
            {
                "page": page,
                "object": name,
                "code": "reduce_overflow_font",
                "minimum_font_pt": min(sizes),
            }
        )
    if changes:
        presentation.save(path)
    return changes


def repair_layout(path, issues):
    changes = repair_overflow(path, issues)
    presentation = Presentation(path)
    width, height = presentation.slide_width, presentation.slide_height
    gap = Pt(8)

    def fits(shape, slide, left, top):
        if left < 0 or top < 0 or left + shape.width > width or top + shape.height > height:
            return False
        for other in slide.shapes:
            if other.shape_id == shape.shape_id or (
                other.width >= width and other.height >= height
            ):
                continue
            overlap_w = min(left + shape.width, other.left + other.width) - max(left, other.left)
            overlap_h = min(top + shape.height, other.top + other.height) - max(top, other.top)
            if overlap_w > Pt(1) and overlap_h > Pt(1):
                return False
        return True

    for issue in issues:
        page = issue.get("page")
        if (
            isinstance(page, bool)
            or not isinstance(page, int)
            or not 1 <= page <= len(presentation.slides)
        ):
            continue
        slide = presentation.slides[page - 1]
        if issue.get("code") == "object_out_of_bounds":
            shape = next((s for s in slide.shapes if s.name == issue.get("object")), None)
            if shape is None or shape.width > width or shape.height > height:
                continue
            left = max(0, min(shape.left, width - shape.width))
            top = max(0, min(shape.top, height - shape.height))
            if fits(shape, slide, left, top):
                shape.left, shape.top = left, top
                changes.append({"page": page, "object": shape.name, "code": "move_inside_slide"})
        elif issue.get("code") == "text_overlap":
            names = issue.get("objects", [])
            if not isinstance(names, list) or len(names) != 2:
                continue
            anchor = next((s for s in slide.shapes if s.name == names[0]), None)
            shape = next((s for s in slide.shapes if s.name == names[1]), None)
            if anchor is None or shape is None or not shape.has_text_frame:
                continue
            positions = [
                (shape.left, anchor.top + anchor.height + gap),
                (shape.left, anchor.top - shape.height - gap),
                (anchor.left + anchor.width + gap, shape.top),
                (anchor.left - shape.width - gap, shape.top),
            ]
            for left, top in positions:
                if fits(shape, slide, left, top):
                    shape.left, shape.top = left, top
                    changes.append(
                        {"page": page, "object": shape.name, "code": "separate_textboxes"}
                    )
                    break
    if any(change["code"] != "reduce_overflow_font" for change in changes):
        presentation.save(path)
    return changes
