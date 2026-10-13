"""Office-independent SVG/PNG previews from resolved editable PPTX geometry."""

import html
import json
import math
from pathlib import Path

from pptx import Presentation


def _rgb(color, default="223A55"):
    try:
        return str(color.rgb) if color.rgb is not None else default
    except (AttributeError, ValueError, TypeError):
        return default


def _fill(fill, default="FFFFFF"):
    try:
        return _rgb(fill.fore_color, default)
    except (AttributeError, ValueError, TypeError):
        return default


def _text(frame, box, name, fill=None):
    lines = []
    for paragraph in frame.paragraphs:
        run = paragraph.runs[0] if paragraph.runs else None
        font = run.font if run and run.font.size else paragraph.font
        size = font.size.pt if font.size else 18
        lines.append(
            {
                "text": paragraph.text,
                "size": size,
                "color": _rgb(font.color),
                "bold": bool(font.bold),
                "font": font.name or "Microsoft YaHei",
            }
        )
    return {
        "kind": "text",
        "id": name,
        **box,
        "lines": lines,
        "margins": [
            frame.margin_left.pt,
            frame.margin_top.pt,
            frame.margin_right.pt,
            frame.margin_bottom.pt,
        ],
        "middle": str(frame.vertical_anchor).startswith("MIDDLE"),
        "fill": fill,
    }


def resolve_pages(pptx_path):
    """Read current geometry, including measured repair changes; never rebuild from stale JSON."""
    presentation = Presentation(pptx_path)
    pages = []
    for number, slide in enumerate(presentation.slides, 1):
        elements = []
        for shape in slide.shapes:
            box = {"x": shape.left.pt, "y": shape.top.pt, "w": shape.width.pt, "h": shape.height.pt}
            if shape.has_table:
                y = box["y"]
                for r, row in enumerate(shape.table.rows):
                    x = box["x"]
                    for c, column in enumerate(shape.table.columns):
                        cell = shape.table.cell(r, c)
                        elements.append(
                            _text(
                                cell.text_frame,
                                {"x": x, "y": y, "w": column.width.pt, "h": row.height.pt},
                                f"{shape.name}:{r + 1}:{c + 1}",
                                _rgb(cell.fill.fore_color, "F1F4F7"),
                            )
                        )
                        elements[-1]["middle"] = str(cell.vertical_anchor).startswith("MIDDLE")
                        x += column.width.pt
                    y += row.height.pt
            elif shape.has_chart:
                chart = shape.chart
                series = [
                    {
                        "name": item.name,
                        "values": list(item.values),
                        "color": _rgb(
                            item.format.line.color
                            if str(chart.chart_type).startswith("LINE")
                            else item.format.fill.fore_color,
                            "0078B8",
                        ),
                    }
                    for item in chart.series
                ]
                elements.append(
                    {
                        "kind": "chart",
                        "id": shape.name,
                        **box,
                        "chart_type": "line"
                        if str(chart.chart_type).startswith("LINE")
                        else "column",
                        "categories": [category.label for category in chart.plots[0].categories],
                        "series": series,
                    }
                )
            elif shape.has_text_frame:
                elements.append(_text(shape.text_frame, box, shape.name))
        pages.append(
            {
                "page": number,
                "width": presentation.slide_width.pt,
                "height": presentation.slide_height.pt,
                "background": _fill(slide.background.fill),
                "elements": elements,
            }
        )
    return pages


def _label(text, x, y, size=14, color="536273", anchor="start", bold=False):
    return (
        f'<text x="{x:.3f}" y="{y:.3f}" font-size="{size}" fill="#{color}" '
        f'text-anchor="{anchor}" font-weight="{"bold" if bold else "normal"}" '
        'font-family="Microsoft YaHei,Noto Sans CJK SC,Arial,sans-serif">'
        f"{html.escape(str(text))}</text>"
    )


def _chart(element):
    x, y, w, h = (element[key] for key in ("x", "y", "w", "h"))
    left, top, width, height = x + 55, y + 15, w - 70, h - 85
    values = [v for item in element["series"] for v in item["values"]]
    low, high = min(0, min(values)), max(0, max(values))
    step = (high - low or 1) / 5
    power = 10 ** math.floor(math.log10(step))
    step = next(factor * power for factor in (1, 2, 5, 10) if factor * power >= step)
    low, high = math.floor(low / step) * step, math.ceil(high / step) * step
    if low == high:
        high = low + step

    def ordinate(value):
        return top + height * (high - value) / (high - low)

    svg = []
    for index in range(round((high - low) / step) + 1):
        value = low + index * step
        yy = ordinate(value)
        svg.append(f'<line x1="{left}" x2="{left + width}" y1="{yy}" y2="{yy}" stroke="#DDE3E8"/>')
        svg.append(_label(f"{value:g}", left - 8, yy + 5, anchor="end"))
    zero = ordinate(0)
    svg.append(f'<line x1="{left}" x2="{left + width}" y1="{zero}" y2="{zero}" stroke="#536273"/>')
    count = len(element["categories"])
    spacing = width / count
    for index, category in enumerate(element["categories"]):
        center = left + spacing * (index + 0.5)
        svg.append(_label(category, center, top + height + 22, anchor="middle"))
    for series_index, series in enumerate(element["series"]):
        if element["chart_type"] == "line":
            points = " ".join(
                f"{left + spacing * (i + 0.5)},{ordinate(v)}"
                for i, v in enumerate(series["values"])
            )
            svg.append(
                f'<polyline points="{points}" fill="none" stroke="#{series["color"]}" stroke-width="2.5"/>'
            )
        else:
            bar_width = spacing * 0.7 / len(element["series"])
            for index, value in enumerate(series["values"]):
                xx = left + spacing * index + spacing * 0.15 + series_index * bar_width
                yy = ordinate(value)
                svg.append(
                    f'<rect x="{xx}" y="{min(yy, zero)}" width="{bar_width}" height="{abs(yy - zero)}" fill="#{series["color"]}"/>'
                )
                svg.append(
                    _label(
                        f"{value:g}",
                        xx + bar_width / 2,
                        yy - 5 if value >= 0 else yy + 17,
                        anchor="middle",
                        color="223A55",
                    )
                )
        legend_x = x + 30 + series_index * w / len(element["series"])
        svg.append(_label(series["name"], legend_x, y + h - 8, color=series["color"]))
    return "".join(svg)


def page_svg(page):
    svg, issues, rectangles = [], [], []
    width, height = page["width"], page["height"]
    svg.append(
        f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 {width} {height}">'
    )
    svg.append(f'<rect width="{width}" height="{height}" fill="#{page["background"]}"/>')
    for element in page["elements"]:
        x, y, w, h = (element[key] for key in ("x", "y", "w", "h"))
        if x < -1 or y < -1 or x + w > width + 1 or y + h > height + 1:
            issues.append(
                {"page": page["page"], "code": "object_out_of_bounds", "object": element["id"]}
            )
        if element["kind"] == "chart":
            svg.append(_chart(element))
            continue
        if element["fill"]:
            svg.append(
                f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#{element["fill"]}"/>'
            )
        ml, mt, mr, mb = element["margins"]
        ink_height = sum(line["size"] * 1.15 for line in element["lines"])
        ink_width = max(
            (
                sum(line["size"] * (1 if ord(c) > 255 else 0.55) for c in line["text"])
                for line in element["lines"]
            ),
            default=0,
        )
        available_h, available_w = h - mt - mb, w - ml - mr
        if any(line["text"] for line in element["lines"]):
            if ink_height > available_h + 1.5 or ink_width > available_w + 1.5:
                issues.append(
                    {
                        "page": page["page"],
                        "code": "text_overflow",
                        "object": element["id"],
                        "bound_height_pt": ink_height,
                        "available_height_pt": available_h,
                        "bound_width_pt": ink_width,
                        "available_width_pt": available_w,
                    }
                )
            ink_y = y + mt + (max(0, available_h - ink_height) / 2 if element["middle"] else 0)
            rectangles.append((element["id"], x + ml, ink_y, ink_width, ink_height))
        baseline = y + mt + (max(0, available_h - ink_height) / 2 if element["middle"] else 0)
        for line in element["lines"]:
            baseline += line["size"]
            svg.append(
                _label(
                    line["text"], x + ml, baseline, line["size"], line["color"], bold=line["bold"]
                )
            )
            baseline += line["size"] * 0.15
    for index, (name, x, y, w, h) in enumerate(rectangles):
        for other, xx, yy, ww, hh in rectangles[index + 1 :]:
            if min(x + w, xx + ww) - max(x, xx) > 2 and min(y + h, yy + hh) - max(y, yy) > 2:
                issues.append(
                    {"page": page["page"], "code": "text_overlap", "objects": [name, other]}
                )
    return "".join(svg) + "</svg>", issues


def render_native(pptx_path, output):
    try:
        import pymupdf
    except ImportError as error:
        from .preview import PreviewUnavailable

        raise PreviewUnavailable(
            "Install ppt-generate[preview] for Office-independent PNG preview"
        ) from error
    output = Path(output)
    pages = resolve_pages(pptx_path)
    report = {
        "backend": "native",
        "pages": [],
        "issues": [],
        "scope": "resolved PPTX geometry, editable text/table/chart data and estimated text bounds",
        "limitations": [
            "PNG/SVG font metrics and chart internals may differ from Office. These images are not actual PowerPoint rendering; human visual/factual review is required."
        ],
        "visual_review": "required",
        "office_verification": "not_performed",
    }
    gallery = [
        '<!doctype html><meta charset="utf-8"><title>Slide preview</title><style>body{background:#edf1f5;font-family:sans-serif;margin:24px}section{max-width:1200px;margin:24px auto}svg{width:100%;height:auto;box-shadow:0 2px 12px #bbc}</style><h1>Slide preview</h1>'
    ]
    for page in pages:
        svg, issues = page_svg(page)
        stem = f"slide-{page['page']:03d}"
        (output / f"{stem}.svg").write_text(svg, encoding="utf-8")
        with pymupdf.open("svg", svg.encode("utf-8")) as document:
            pdf = document.convert_to_pdf()
        with pymupdf.open("pdf", pdf) as document:
            document[0].get_pixmap(
                matrix=pymupdf.Matrix(1600 / document[0].rect.width, 1600 / document[0].rect.width)
            ).save(output / f"{stem}.png")
        report["pages"].append({"page": page["page"], "file": f"{stem}.png", "svg": f"{stem}.svg"})
        report["issues"].extend(issues)
        gallery.append(f"<section><h2>Slide {page['page']}</h2>{svg}</section>")
    (output / "index.html").write_text("".join(gallery), encoding="utf-8")
    (output / "scene.json").write_text(
        json.dumps(pages, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    report["status"] = "issues_found" if report["issues"] else "passed"
    return report
