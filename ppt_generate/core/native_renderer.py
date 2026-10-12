"""Deterministic editable PPTX renderer. Layout is owned by code, not the LLM."""

import unicodedata

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_DATA_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .deck import validate_deck
from .editorial import review_editorial

INK = "223A55"
BLUE = "0078B8"
GRAY = "536273"


def _font(font, size, color=None, bold=False):
    font.name = "Microsoft YaHei"
    font.size = Pt(size)
    font.bold = bold
    if color:
        font.color.rgb = RGBColor.from_string(color)
    # Set East Asian typeface explicitly rather than inheriting a serif theme font.
    east_asian = font._rPr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        font._rPr.append(east_asian)
    east_asian.set("typeface", "Microsoft YaHei")


def _wrap(text, width):
    lines, line, used = [], "", 0
    for char in text:
        size = 2 if unicodedata.east_asian_width(char) in "WF" else 1
        if char == "\n":
            lines.append(line)
            line, used = "", 0
        elif used + size > width:
            space = line.rfind(" ")
            tail = line[space + 1 :] if space >= 0 else ""
            tail_width = sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in tail)
            if space > 0 and tail_width < width * 0.35:
                lines.append(line[:space].rstrip())
                line, used = tail, tail_width
            else:
                lines.append(line)
                line, used = "", 0
        if char != "\n":
            line += char
            used += size
    lines.append(line)
    return "\n".join(lines)


def _box(slide, text, x, y, w, h, size=22, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    # Conservative wrapping avoids shrinking body text to fit dense content.
    wrapped = _wrap(text, max(8, int(w * 72 / size * 1.65)))
    required = len(wrapped.splitlines()) * size * 1.35 / 72 + 0.04
    if required > h:
        raise ValueError("Text exceeds layout capacity; shorten or split the page")
    for i, line in enumerate(wrapped.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        _font(p.font, size, color, bold)
        p.line_spacing = 1.15
        p.space_after = Pt(0)
    return shape


def build_presentation(deck):
    validate_deck(deck)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for number, spec in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        layout = spec["layout"]
        if layout == "cover":
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(INK)
            _box(slide, spec["title"], 0.8, 2, 11.7, 2, 40, "FFFFFF", True)
            if spec.get("subtitle"):
                _box(slide, spec["subtitle"], 0.8, 4.4, 11.7, 1.5, 22, "DDE6EF")
        else:
            _box(slide, spec["title"], 0.6, 0.35, 12.1, 1.15, 30, bold=True)
            if layout == "kpi":
                count = len(spec["metrics"])
                width = (11.9 - 0.35 * (count - 1)) / count
                for index, metric in enumerate(spec["metrics"]):
                    x = 0.7 + index * (width + 0.35)
                    _box(slide, metric["label"], x, 1.8, width, 1.0, 22, GRAY)
                    _box(slide, metric["value"], x, 3.0, width, 1.45, 34, BLUE, True)
                    if metric.get("detail"):
                        _box(slide, metric["detail"], x, 4.75, width, 1.55, 18)
            elif layout == "timeline":
                height = 4.8 / len(spec["events"])
                for index, event in enumerate(spec["events"]):
                    y = 1.65 + index * height
                    _box(slide, event["date"], 0.7, y, 2.1, height - 0.1, 20, BLUE, True)
                    _box(slide, event["heading"], 3.0, y, 3.1, height - 0.1, 20, INK, True)
                    _box(slide, event["detail"], 6.4, y, 6.1, height - 0.1, 18)
            elif layout == "bullets":
                y = 1.7
                size = 26 if len(spec["bullets"]) <= 4 else 22
                for item in spec["bullets"]:
                    wrapped = _wrap(item, int(11.7 * 72 / size * 1.65) - 2)
                    height = len(wrapped.splitlines()) * size * 1.35 / 72 + 0.12
                    if y + height > 6.5:
                        raise ValueError(f"Slide {number}: too much bullet content")
                    _box(slide, "• " + wrapped, 0.8, y, 11.7, height, size)
                    y += height + 0.2
            elif layout == "split":
                for index, column in enumerate(spec["columns"]):
                    x = 0.8 + index * 6.1
                    _box(slide, column["heading"], x, 1.7, 5.6, 0.8, 24, BLUE, True)
                    y = 2.7
                    for item in column["bullets"]:
                        text = _wrap(item, int(5.6 * 72 / 21 * 1.65))
                        height = len(text.splitlines()) * 21 * 1.35 / 72 + 0.12
                        if y + height > 6.5:
                            raise ValueError(f"Slide {number}: column content exceeds capacity")
                        _box(slide, text, x, y, 5.6, height, 21)
                        y += height + 0.16
            elif layout == "table":
                rows = [spec["headers"], *spec["rows"]]
                table = slide.shapes.add_table(
                    len(rows), len(rows[0]), Inches(0.7), Inches(1.7), Inches(11.9), Inches(4.7)
                ).table
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        cell = table.cell(r, c)
                        width = 11.9 / len(row)
                        text = _wrap(value, int(width * 72 / 18 * 1.6))
                        if len(text.splitlines()) * 18 * 1.35 / 72 > 4.7 / len(rows) - 0.15:
                            raise ValueError(f"Slide {number}: table cell content exceeds capacity")
                        cell.text = text
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        for p in cell.text_frame.paragraphs:
                            _font(p.font, 18, "FFFFFF" if r == 0 else INK, r == 0)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(INK if r == 0 else "F1F4F7")
            elif layout == "chart":
                unit_label = (
                    "单位"
                    if any("\u4e00" <= char <= "\u9fff" for char in spec["title"])
                    else "Unit"
                )
                _box(
                    slide,
                    f"{spec['period']}    {unit_label}: {spec['unit']}",
                    0.8,
                    1.5,
                    11.7,
                    0.5,
                    18,
                    GRAY,
                )
                data = CategoryChartData()
                data.categories = spec["categories"]
                for series in spec["series"]:
                    data.add_series(series["name"], series["values"])
                kind = (
                    XL_CHART_TYPE.LINE
                    if spec.get("chart_type") == "line"
                    else XL_CHART_TYPE.COLUMN_CLUSTERED
                )
                chart = slide.shapes.add_chart(
                    kind, Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.3), data
                ).chart
                chart.has_legend = True
                chart.has_title = False
                chart.plots[0].vary_by_categories = False
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                # python-pptx maps this flag to c:overlay; False reserves legend space.
                chart.legend.include_in_layout = False
                _font(chart.legend.font, 14, GRAY)
                _font(chart.category_axis.tick_labels.font, 14, GRAY)
                chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
                _font(chart.value_axis.tick_labels.font, 14, GRAY)
                if all(value >= 0 for series in spec["series"] for value in series["values"]):
                    chart.value_axis.minimum_scale = 0
                chart.value_axis.has_major_gridlines = True
                chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(
                    "DDE3E8"
                )
                chart.value_axis.major_gridlines.format.line.width = Pt(0.6)
                for index, series in enumerate(chart.series):
                    color = [BLUE, INK, GRAY][index]
                    if spec.get("chart_type") == "line":
                        series.format.line.color.rgb = RGBColor.from_string(color)
                        series.format.line.width = Pt(2.5)
                    else:
                        series.invert_if_negative = False
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
                if spec.get("chart_type") != "line":
                    chart.plots[0].has_data_labels = True
                    chart.plots[0].data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
                    _font(chart.plots[0].data_labels.font, 14, INK)
                chart.value_axis.has_title = True
                chart.value_axis.axis_title.text_frame.text = spec["unit"]
                for p in chart.value_axis.axis_title.text_frame.paragraphs:
                    _font(p.font, 14, GRAY)
        footer_color = "DDE6EF" if layout == "cover" else GRAY
        _box(slide, spec.get("source", ""), 0.7, 6.85, 11.6, 0.45, 12, footer_color)
        _box(slide, str(number), 12.25, 6.85, 0.5, 0.4, 12, footer_color)
        slide.notes_slide.notes_text_frame.text = (
            spec.get("notes", "") + "\n" + spec.get("source", "")
        )
    return prs


def render_deck(deck, output):
    prs = build_presentation(deck)
    prs.save(str(output))
    return {
        "status": "structural_checks_passed",
        "checks": ["schema", "conservative_text_capacity", "editable_objects"],
        "slide_count": len(prs.slides),
        "editorial": review_editorial(deck),
        "visual_review": "required",
        "limitations": [
            "Text capacity uses conservative estimates, not font rasterization; render review is still required."
        ],
    }
